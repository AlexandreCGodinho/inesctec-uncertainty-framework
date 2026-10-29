#!/usr/bin/env python3
"""
add_helper_columns_finqa.py

Automates the addition of helper columns ('value' and 'unit') to the FinQA labeling grid,
as requested by Prof. Alvaro to enable inter-annotator agreement computation under Doc 07 tolerance.

Columns structure (25 columns total):
- 1..5: query_id, Pattern A, Pattern B, Pattern C, Pattern D
- For each pattern X in (A, B, C, D):
    intended_number_X, abstain_X, abstain_reason_X, value_X, unit_X
"""

import argparse
import csv
import re
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parent.parent


def parse_label(raw: Any, abst: Any) -> tuple:
    """
    Extracts plain numeric value (with dot decimal) and standardized unit from free-form label.
    Follows Doc 07 rules (strips currency/thousands separators, identifies scale/unit).
    """
    if str(abst).strip().lower() == "true" or raw is None:
        return None, None

    s = str(raw).strip()
    if not s or s.lower() in ("none", "null"):
        return None, None
    if s.lower() in ("yes", "no"):
        return s.lower(), "boolean"
    if s.lower() == "inconclusive":
        return "inconclusive", "text"

    clean = s
    unit = ""

    if "%" in s:
        unit = "%"
        clean = clean.replace("%", "").strip()
    elif "square feet" in s.lower() or "sqft" in s.lower():
        unit = "square feet"
        clean = re.sub(r"square feet|sqft", "", clean, flags=re.I).strip()
    elif "billion" in s.lower():
        unit = "billion"
        clean = re.sub(r"billion", "", clean, flags=re.I).strip()
    elif "million" in s.lower():
        unit = "million"
        clean = re.sub(r"million", "", clean, flags=re.I).strip()
    elif "thousand" in s.lower() or "$000s" in s.lower() or "000s" in s.lower():
        unit = "thousand"
        clean = re.sub(
            r"\(in\s*\$000s\)|in\s*\$000s|\(in\s*thousands\)|in\s*thousands|thousand",
            "",
            clean,
            flags=re.I,
        ).strip()
    elif "$" in s:
        unit = "$"
        clean = clean.replace("$", "").strip()
    else:
        unit = ""

    # Remove currency symbol, parentheses and thousands separator commas
    clean = clean.replace("$", "").replace("(", "").replace(")", "").replace(",", "").strip()
    return clean, unit


def _read_input_rows(input_path: Path) -> list[dict[str, Any]]:
    if input_path.suffix.lower() == ".csv":
        with open(input_path, "r", encoding="utf-8") as f:
            return list(csv.DictReader(f))

    import openpyxl

    wb_orig = openpyxl.load_workbook(str(input_path), data_only=True)
    ws_orig = wb_orig["FinQA Labeling"] if "FinQA Labeling" in wb_orig.sheetnames else wb_orig.active
    headers = [ws_orig.cell(1, c).value for c in range(1, ws_orig.max_column + 1)]

    # Support both named headers and legacy 17-column positional layout
    has_named = "intended_number_A" in [str(h or "") for h in headers]
    orig_map = {"A": (6, 7, 8), "B": (9, 10, 11), "C": (12, 13, 14), "D": (15, 16, 17)}

    rows: list[dict[str, Any]] = []
    for r in range(2, ws_orig.max_row + 1):
        if has_named:
            row_dict = {
                str(headers[c - 1]): ws_orig.cell(r, c).value
                for c in range(1, ws_orig.max_column + 1)
                if headers[c - 1] is not None
            }
        else:
            row_dict = {
                "query_id": ws_orig.cell(r, 1).value,
                "Pattern A": ws_orig.cell(r, 2).value,
                "Pattern B": ws_orig.cell(r, 3).value,
                "Pattern C": ws_orig.cell(r, 4).value,
                "Pattern D": ws_orig.cell(r, 5).value,
            }
            for pat, (c_num, c_abs, c_rea) in orig_map.items():
                row_dict[f"intended_number_{pat}"] = ws_orig.cell(r, c_num).value
                row_dict[f"abstain_{pat}"] = ws_orig.cell(r, c_abs).value
                row_dict[f"abstain_reason_{pat}"] = ws_orig.cell(r, c_rea).value
        rows.append(row_dict)
    return rows


def process_grid(input_path: str, output_path: str) -> None:
    import openpyxl

    in_p = Path(input_path)
    out_p = Path(output_path)
    print(f"Loading input file: {in_p}")
    rows = _read_input_rows(in_p)

    wb_new = openpyxl.Workbook()
    ws_new = wb_new.active
    ws_new.title = "FinQA Labeling"

    new_headers = [
        "query_id", "Pattern A", "Pattern B", "Pattern C", "Pattern D",
        "intended_number_A", "abstain_A", "abstain_reason_A", "value_A", "unit_A",
        "intended_number_B", "abstain_B", "abstain_reason_B", "value_B", "unit_B",
        "intended_number_C", "abstain_C", "abstain_reason_C", "value_C", "unit_C",
        "intended_number_D", "abstain_D", "abstain_reason_D", "value_D", "unit_D",
    ]
    ws_new.append(new_headers)

    for row in rows:
        row_data = [
            row.get("query_id"),
            row.get("Pattern A"),
            row.get("Pattern B"),
            row.get("Pattern C"),
            row.get("Pattern D"),
        ]
        for pat in ["A", "B", "C", "D"]:
            num_v = row.get(f"intended_number_{pat}")
            abs_v = row.get(f"abstain_{pat}")
            rea_v = row.get(f"abstain_reason_{pat}")
            val_v, unit_v = parse_label(num_v, abs_v)
            row_data.extend([num_v, abs_v, rea_v, val_v, unit_v])
        ws_new.append(row_data)

    out_p.parent.mkdir(parents=True, exist_ok=True)
    wb_new.save(str(out_p))
    print(f"Successfully written {ws_new.max_row} rows and {ws_new.max_column} columns to: {out_p}")


def main() -> None:
    default_dir = PROJECT_ROOT / "data" / "finqa_labeling"
    parser = argparse.ArgumentParser(
        description="Add value/unit helper columns to a FinQA labeling grid (.xlsx or .csv)."
    )
    parser.add_argument(
        "--input",
        type=str,
        default=str(default_dir / "finqa_labeling_grid_fixed.csv"),
        help="Input FinQA labeling grid (.xlsx or .csv)",
    )
    parser.add_argument(
        "--output",
        type=str,
        default=str(default_dir / "finqa_labeling_grid_fixed_helper_columns_AG.xlsx"),
        help="Output .xlsx path with helper columns",
    )
    args = parser.parse_args()

    in_file = Path(args.input)
    if not in_file.exists():
        raise FileNotFoundError(f"Input file not found at: {in_file}")
    process_grid(str(in_file), str(args.output))


if __name__ == "__main__":
    main()
